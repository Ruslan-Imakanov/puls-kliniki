"""Шаблоны для продаж: счёт + акт и таблица учёта обращений. python3 build_templates.py
Синие ячейки — вводить, чёрные — формулы."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

BLUE = Font(name='Calibri', size=11, color='0000FF'); BOLD = Font(name='Calibri', size=11, bold=True)
H1 = Font(name='Calibri', size=16, bold=True); GREY = Font(name='Calibri', size=10, color='666666', italic=True)
HEAD = PatternFill('solid', fgColor='1F6F66'); HF = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
INP = PatternFill('solid', fgColor='EAF3FF'); thin = Side(style='thin', color='BBBBBB'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0.00" ₸"'

wb = Workbook()
# ---------- Реквизиты ----------
r = wb.active; r.title = 'Реквизиты'
r['A1'] = 'Реквизиты и данные клиента'; r['A1'].font = H1
r['A2'] = 'Заполните синие ячейки один раз — счёт и акт подставят данные сами.'; r['A2'].font = GREY
rows = [
 ('Исполнитель', None),
 ('Наименование', 'ИП Имаканов Руслан Арсланович'), ('ИИН', '020413550395'), ('Адрес', 'г. Костанай, ул. Тауелсиздик, 70'),
 ('Банк', 'АО «Народный Банк Казахстана»'), ('ИИК', 'KZ56601A221039793271'), ('БИК', 'HSBKKZKX'), ('Кбе', '19'),
 ('Телефон', '+7 707 611 30 75'), ('Почта', 'ruslanimakanov8@gmail.com'),
 ('Заказчик (клиника)', None),
 ('Наименование', ''), ('БИН/ИИН', ''), ('Адрес', ''),
 ('Документ', None),
 ('Номер счёта', 1), ('Дата счёта', None), ('Номер акта', 1), ('Дата акта', None), ('Договор', 'Публичная оферта https://puls-kliniki.imakanovr.workers.dev/terms'),
]
row = 4
pos = {}
for k, v in rows:
    if v is None and k not in ('Дата счёта', 'Дата акта'):
        r.cell(row, 1, k).font = BOLD; row += 1; continue
    r.cell(row, 1, k); c = r.cell(row, 2, v); c.font = BLUE; c.fill = INP; c.border = BOX
    pos[(len([1 for kk in pos if kk[0] == k]), k)] = row  # не используется, оставлено для читаемости
    if k in ('Дата счёта', 'Дата акта'): c.value = '=TODAY()'; c.number_format = 'DD.MM.YYYY'
    row += 1
r.column_dimensions['A'].width = 24; r.column_dimensions['B'].width = 60
# адреса ячеек (фиксированы порядком выше)
R = dict(ip=5, iin=6, addr=7, bank=8, iik=9, bik=10, kbe=11, tel=12, mail=13, cl=15, cl_bin=16, cl_addr=17, n_inv=19, d_inv=20, n_act=21, d_act=22, dog=23)
r['A25'] = 'Тарифы (справочник)'; r['A25'].font = BOLD
for i, (n, p) in enumerate([('Подписка «Старт», месяц', 29000), ('Подписка «Клиника», месяц', 49000), ('Подписка «Сеть», месяц', 89000), ('Подключение (разово)', 150000)]):
    r.cell(26 + i, 1, n); c = r.cell(26 + i, 2, p); c.font = BLUE; c.fill = INP; c.number_format = '#,##0" ₸"'; c.border = BOX

def doc(ws, title, num_key, date_key, who_a, who_b):
    ws['A1'] = f'=\"{title} № \"&Реквизиты!B{R[num_key]}&\" от \"&TEXT(Реквизиты!B{R[date_key]},\"DD.MM.YYYY\")'; ws['A1'].font = H1
    ws['A3'] = who_a; ws['A3'].font = BOLD
    ws['B3'] = f'=Реквизиты!B{R["ip"]}&\", ИИН \"&Реквизиты!B{R["iin"]}&\", \"&Реквизиты!B{R["addr"]}'
    ws['A4'] = who_b; ws['A4'].font = BOLD
    ws['B4'] = f'=Реквизиты!B{R["cl"]}&IF(Реквизиты!B{R["cl_bin"]}<>\"\",\", БИН/ИИН \"&Реквизиты!B{R["cl_bin"]},\"\")&IF(Реквизиты!B{R["cl_addr"]}<>\"\",\", \"&Реквизиты!B{R["cl_addr"]},\"\")'
    ws['A5'] = 'Договор'; ws['A5'].font = BOLD; ws['B5'] = f'=Реквизиты!B{R["dog"]}'
    for j, h in enumerate(['№', 'Наименование', 'Кол-во', 'Цена', 'Сумма']):
        c = ws.cell(7, 1 + j, h); c.font = HF; c.fill = HEAD; c.alignment = Alignment(horizontal='center'); c.border = BOX
    items = [('Подписка на программу «Пульс клиники», тариф «Клиника», 1 месяц', 1, '=Реквизиты!B27'),
             ('Подключение: перенос данных, настройка, обучение', 1, '=Реквизиты!B29'), ('', None, None), ('', None, None), ('', None, None)]
    for i, (n, q, p) in enumerate(items):
        rr = 8 + i
        ws.cell(rr, 1, i + 1)
        c = ws.cell(rr, 2, n); c.font = BLUE; c.fill = INP
        c = ws.cell(rr, 3, q); c.font = BLUE; c.fill = INP
        c = ws.cell(rr, 4, p); c.font = BLUE; c.fill = INP; c.number_format = MONEY
        ws.cell(rr, 5, f'=IF(AND(ISNUMBER(C{rr}),ISNUMBER(D{rr})),C{rr}*D{rr},0)').number_format = MONEY
        for cc in range(1, 6): ws.cell(rr, cc).border = BOX
    ws['D13'] = 'Итого'; ws['D13'].font = BOLD; ws['E13'] = '=SUM(E8:E12)'; ws['E13'].number_format = MONEY; ws['E13'].font = BOLD
    ws['D14'] = 'НДС'; ws['E14'] = 'без НДС'; ws['E14'].alignment = Alignment(horizontal='right')
    ws['A16'] = 'Всего к оплате:'; ws['A16'].font = BOLD
    ws['B16'] = '=TEXT(E13,"#,##0")&" тенге"'
    ws['A17'] = 'Сумма прописью:'; ws['A17'].font = BOLD
    c = ws['B17']; c.value = 'введите вручную, например: сто девяносто девять тысяч тенге 00 тиын'; c.font = BLUE; c.fill = INP
    ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 62; ws.column_dimensions['C'].width = 10; ws.column_dimensions['D'].width = 18; ws.column_dimensions['E'].width = 18
    ws.page_setup.orientation = 'portrait'; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 1

s = wb.create_sheet('Счёт'); doc(s, 'Счёт на оплату', 'n_inv', 'd_inv', 'Поставщик', 'Покупатель')
s['A19'] = 'Реквизиты для оплаты'; s['A19'].font = BOLD
pay = [('Получатель', f'=Реквизиты!B{R["ip"]}'), ('ИИН', f'=Реквизиты!B{R["iin"]}'), ('Банк', f'=Реквизиты!B{R["bank"]}'),
       ('ИИК', f'=Реквизиты!B{R["iik"]}'), ('БИК', f'=Реквизиты!B{R["bik"]}'), ('Кбе', f'=Реквизиты!B{R["kbe"]}'),
       ('Назначение платежа', f'="Оплата по счёту № "&Реквизиты!B{R["n_inv"]}&" от "&TEXT(Реквизиты!B{R["d_inv"]},"DD.MM.YYYY")&", подписка «Пульс клиники»"')]
for i, (k, f) in enumerate(pay): s.cell(20 + i, 1, k); s.cell(20 + i, 2, f)
s['A28'] = 'Счёт действителен 10 календарных дней. Подпись: ____________  / Имаканов Р. А. /'; s['A28'].font = GREY

a = wb.create_sheet('Акт'); doc(a, 'Акт выполненных работ (оказанных услуг)', 'n_act', 'd_act', 'Исполнитель', 'Заказчик')
a['A19'] = 'Услуги оказаны в полном объёме и в срок. Заказчик претензий по объёму, качеству и срокам не имеет.'
a['A21'] = 'Исполнитель: ____________  / Имаканов Р. А. /'; a['A22'] = 'Заказчик: ____________  / _______________ /'

# ---------- Учёт обращений ----------
t = wb.create_sheet('Учёт обращений')
head = ['Клиника', 'Город', 'Страна', 'Контакт (имя, должность)', 'Почта / телефон', 'Канал', 'Дата 1-го письма', 'Метка (utm)', 'Ответ', 'Статус', 'Следующий шаг', 'Когда', 'Тариф', 'Сумма, ₸', 'Дней с письма', 'Нужно напомнить?']
for j, h in enumerate(head):
    c = t.cell(1, 1 + j, h); c.font = HF; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BOX
sample = [('MEDIKER', 'Шымкент', 'Казахстан', '', 'pochta@mediker.kz', 'Почта', '2026-10-24', 'kz1', '', 'Написали', 'Напомнить', None, '', None),
          ('Эмирмед', 'Астана', 'Казахстан', '', 'emirmed@mail.ru', 'Почта', '2026-10-24', 'kz2', '', 'Написали', 'Напомнить', None, '', None),
          ('AKFA Medline', 'Ташкент', 'Узбекистан', '', 'info@akfamedline.uz', 'Почта', '2026-10-24', 'uz1', '', 'Написали', 'Напомнить', None, '', None)]
import datetime
for i in range(200):
    rr = 2 + i
    if i < len(sample):
        vals = sample[i]
        for j, v in enumerate(vals):
            if j == 6: v = datetime.date.fromisoformat(v)
            if j == 11: v = datetime.date(2026, 10, 27)
            t.cell(rr, 1 + j, v)
    for j in range(13):
        c = t.cell(rr, 1 + j); c.font = BLUE; c.border = BOX
    t.cell(rr, 7).number_format = 'DD.MM.YYYY'; t.cell(rr, 12).number_format = 'DD.MM.YYYY'; t.cell(rr, 14).number_format = '#,##0'
    t.cell(rr, 15, f'=IF(G{rr}="","",TODAY()-G{rr})')
    t.cell(rr, 16, f'=IF(OR(G{rr}="",J{rr}="Оплатили",J{rr}="Отказ",J{rr}="Ответили"),"",IF(TODAY()-G{rr}>=3,"да","")))')
    t.cell(rr, 16).value = f'=IF(OR(G{rr}="",J{rr}="Оплатили",J{rr}="Отказ",J{rr}="Показ",J{rr}="Пробный месяц"),"",IF(TODAY()-G{rr}>=3,"да",""))'
    for j in (15, 16): t.cell(rr, j).border = BOX
dv = DataValidation(type='list', formula1='"Написали,Ответили,Показ,Пробный месяц,Оплатили,Отказ"', allow_blank=True); t.add_data_validation(dv); dv.add('J2:J201')
dv2 = DataValidation(type='list', formula1='"Старт,Клиника,Сеть"', allow_blank=True); t.add_data_validation(dv2); dv2.add('M2:M201')
t.conditional_formatting.add('P2:P201', CellIsRule(operator='equal', formula=['"да"'], fill=PatternFill('solid', bgColor='FFE0B2')))
t.conditional_formatting.add('J2:J201', CellIsRule(operator='equal', formula=['"Оплатили"'], fill=PatternFill('solid', bgColor='C8E6C9')))
t.conditional_formatting.add('J2:J201', CellIsRule(operator='equal', formula=['"Отказ"'], fill=PatternFill('solid', bgColor='EEEEEE')))
for col, w in zip('ABCDEFGHIJKLMNOP', [24, 14, 13, 24, 28, 10, 14, 10, 30, 14, 24, 13, 10, 12, 10, 14]): t.column_dimensions[col].width = w
t.row_dimensions[1].height = 34; t.freeze_panes = 'B2'; t.auto_filter.ref = 'A1:P201'

# ---------- Итоги воронки ----------
w = wb.create_sheet('Воронка')
w['A1'] = 'Воронка продаж'; w['A1'].font = H1
w['A2'] = 'Считается из листа «Учёт обращений».'; w['A2'].font = GREY
stages = ['Написали', 'Ответили', 'Показ', 'Пробный месяц', 'Оплатили', 'Отказ']
for j, h in enumerate(['Статус', 'Клиник', 'Доля от написанных']):
    c = w.cell(4, 1 + j, h); c.font = HF; c.fill = HEAD
for i, sname in enumerate(stages):
    w.cell(5 + i, 1, sname); w.cell(5 + i, 2, f"=COUNTIF('Учёт обращений'!J2:J201,A{5+i})")
    w.cell(5 + i, 3, f'=IF($B$11=0,0,B{5+i}/$B$11)').number_format = '0%'
w['A11'] = 'Всего'; w['A11'].font = BOLD; w['B11'] = '=SUM(B5:B10)'
w['A13'] = 'Получено, ₸'; w['B13'] = "=SUMIF('Учёт обращений'!J2:J201,\"Оплатили\",'Учёт обращений'!N2:N201)"; w['B13'].number_format = '#,##0'
w['A14'] = 'Нужно напомнить сегодня'; w['B14'] = "=COUNTIF('Учёт обращений'!P2:P201,\"да\")"
w['A15'] = 'Из написанных ответили (%)'; w['B15'] = '=IF(B11=0,0,(B6+B7+B8+B9+B10)/B11)'; w['B15'].number_format = '0%'
w.column_dimensions['A'].width = 30; w.column_dimensions['B'].width = 14; w.column_dimensions['C'].width = 20

wb._sheets = [wb['Реквизиты'], wb['Счёт'], wb['Акт'], wb['Учёт обращений'], wb['Воронка']]
wb.save('Счёт_акт_и_учёт_обращений.xlsx'); print('ok')
